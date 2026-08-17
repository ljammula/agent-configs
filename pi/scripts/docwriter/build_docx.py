"""Build a deterministic styled DOCX from a small Markdown subset."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)


def set_font(font, name="Calibri", size=11, color=None, bold=False):
    font.name = name
    font.size = Pt(size)
    font.bold = bold
    if color:
        font.color.rgb = color


def add_inline(paragraph, text: str, size=11, bold=False):
    """Render the small inline Markdown subset used by the draft prompts."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_font(run.font, size=size, bold=bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run.font, size=size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run.font, size=size)
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run.font, "Courier New", size=size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run.font, size=size, bold=bold)


def add_rich_paragraph(document, text: str, style=None):
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text)
    return paragraph


def cell_margins(cell):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if margins.getparent() is None:
        props.append(margins)
    for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        node = margins.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
        if node.getparent() is None:
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure(document):
    section = document.sections[0]
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, attr, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    set_font(normal.font)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = document.styles[name]
        set_font(style.font, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        set_font(style.font)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.style.font, size=9, color=RGBColor(0x66, 0x66, 0x66))
    footer.add_run("Page ")
    run = footer.add_run()
    for tag, value in (("w:fldChar", "begin"), ("w:instrText", " PAGE "), ("w:fldChar", "separate"), ("w:t", "1"), ("w:fldChar", "end")):
        node = OxmlElement(tag)
        if tag == "w:instrText":
            node.set(qn("xml:space"), "preserve")
        if tag == "w:fldChar":
            node.set(qn("w:fldCharType"), value)
        node.text = value if tag == "w:instrText" or tag == "w:t" else None
        run._r.append(node)


def add_table(document, rows):
    count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=count)
    table.style = "Table Grid"
    table.autofit = False
    widths = [1.875, 4.625] if count == 2 else [6.5 / count] * count
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            row_properties = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row_properties.append(header)
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            cell.text = ""
            if row_index == 0:
                props = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F4F7")
                props.append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, rows[row_index][col_index].strip(), size=10, bold=row_index == 0)
    document.add_paragraph()


def render(document, markdown):
    lines = markdown.splitlines()
    paragraph = []
    title_done = False
    index = 0

    def flush():
        nonlocal paragraph
        if paragraph:
            add_rich_paragraph(document, " ".join(x.strip() for x in paragraph))
            paragraph = []

    while index < len(lines):
        line = lines[index].rstrip()
        if line.strip() in {"---", "***", "___"}:
            flush()
            index += 1
            continue
        if not line.strip():
            flush()
            index += 1
            continue
        if line.strip().startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$", lines[index + 1]):
            flush()
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([part.strip() for part in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, [rows[0], *rows[2:]])
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush()
            level, text = len(heading.group(1)), heading.group(2).strip()
            if level == 1 and not title_done:
                p = document.add_paragraph(style="Title")
                set_font(p.style.font, size=24, color=DARK_BLUE, bold=True)
                p.paragraph_format.space_after = Pt(12)
                p.add_run(text)
                title_done = True
            else:
                document.add_paragraph(text, style=f"Heading {level}")
            index += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        number = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet or number:
            flush()
            item = (bullet or number).group(1)
            list_paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(list_paragraph, item)
            index += 1
            continue
        if line.lstrip().startswith("> "):
            paragraph.append(line.lstrip()[2:])
            index += 1
            continue
        paragraph.append(line)
        index += 1
    flush()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_markdown", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()
    document = Document()
    configure(document)
    render(document, args.input_markdown.read_text(encoding="utf-8"))
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)
    print(args.output_docx)


if __name__ == "__main__":
    main()
