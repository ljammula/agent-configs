"""Extract supported source files into JSON for the TypeScript pipeline."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SUPPORTED = {".docx", ".md", ".pdf", ".txt"}


def extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n\n".join((page.extract_text() or "").strip() for page in PdfReader(str(path)).pages).strip()


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    raise ValueError(f"unsupported source format: {path.suffix}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+")
    args = parser.parse_args()
    results = []
    errors = []
    for raw in args.paths:
        root = Path(raw)
        candidates = sorted(root.rglob("*") if root.is_dir() else [root])
        for path in candidates:
            if not path.is_file() or path.suffix.lower() not in SUPPORTED:
                continue
            try:
                text = extract(path).strip()
            except Exception as exc:
                errors.append(f"failed to extract {path}: {exc}")
                continue
            if text:
                results.append({"path": str(path), "text": text})
    if not results:
        for error in errors:
            print(error, file=sys.stderr)
        print("no supported, non-empty source files found", file=sys.stderr)
        return 2
    for error in errors:
        print(error, file=sys.stderr)
    json.dump(results, sys.stdout, ensure_ascii=False)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
